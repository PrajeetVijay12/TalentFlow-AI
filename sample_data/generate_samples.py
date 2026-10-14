"""Creates the sample resumes (PDF + DOCX) used for demos and tests.

Run from the project root:   python sample_data/generate_samples.py

Each candidate is designed to exercise a different path through the pipeline.
"""
from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

OUT = Path(__file__).parent / "resumes"

CANDIDATES = [
    # 1. Excellent candidate: everything matches, major project, AWS certificate
    dict(file="rahul_sharma.pdf", name="Rahul Sharma", email="rahul.sharma@example.com", phone="+91 98765 43210",
         location="Bengaluru", edu="B.Tech in Computer Science, NIT Surathkal, 2021", cgpa="CGPA: 8.6/10",
         skills=["Python", "Java", "SQL", "Docker", "Git"], exp=["Software Engineer, Infosys, Jan 2021 - Dec 2023"],
         projects=["Hospital Management System - Major Project: full-stack hospital app with billing",
                   "Online Library - Mini Project: book lending tracker"],
         certs=["AWS Certified Solutions Architect"]),
    # 2. Good candidate: college major project, different certificate
    dict(file="priya_nair.docx", name="Priya Nair", email="priya.nair@example.com", phone="9123456780",
         location="Bengaluru, Karnataka", edu="Bachelor of Technology, VIT Vellore", cgpa="CGPA: 8.1",
         skills=["python", "JAVA", "Sql", "React"], exp=["2 years of experience as a backend developer"],
         projects=["Smart Attendance System - College Major Project: face-recognition free attendance tracker"],
         certs=["Azure Fundamentals"]),
    # 3. Missing required skill (no SQL)
    dict(file="arjun_mehta.pdf", name="Arjun Mehta", email="arjun.mehta@example.com", phone="9988776655",
         location="Bengaluru", edu="BTech Information Technology, SRM University", cgpa="CGPA: 8.0",
         skills=["Python", "Java", "Linux"], exp=["3 years experience in application development"],
         projects=["E-Commerce Portal - Major Project: shopping cart and payments"],
         certs=["AWS Certified Developer"]),
    # 4. Insufficient experience (6 months)
    dict(file="neha_verma.docx", name="Neha Verma", email="neha.verma@example.com", phone="9012345678",
         location="Bengaluru", edu="B.Tech Computer Science, PES University", cgpa="CGPA: 8.4",
         skills=["Python", "Java", "SQL", "Flask"], exp=["6 months of experience as an intern"],
         projects=["Weather Dashboard - Mini Project: API based weather app"],
         certs=["AWS Cloud Practitioner"]),
    # 5. Fails degree and CGPA
    dict(file="karan_singh.pdf", name="Karan Singh", email="karan.singh@example.com", phone="9876501234",
         location="Bengaluru", edu="BCA, Christ University", cgpa="CGPA: 6.4",
         skills=["Python", "Java", "SQL"], exp=["4 years experience in support engineering"],
         projects=["Inventory Tracker - Mini Project"], certs=[]),
    # 6. Eligible but weaker profile (mini project, no certificate) -> wait list
    dict(file="sneha_iyer.docx", name="Sneha Iyer", email="sneha.iyer@example.com", phone="9345678123",
         location="Bengaluru", edu="B.Tech Electronics, RVCE", cgpa="CGPA: 7.6",
         skills=["Python", "Java", "SQL"], exp=["Junior Developer, TCS, Jul 2022 - Aug 2023"],
         projects=["Student Portal - Mini Project: marks and attendance viewer"], certs=[]),
    # 7. Passes BFS/DFS but violates the location constraint (CSP catches it)
    dict(file="vikram_rao.pdf", name="Vikram Rao", email="vikram.rao@example.com", phone="9871234560",
         location="Hyderabad", edu="B.Tech Computer Science, JNTU", cgpa="CGPA: 8.8",
         skills=["Python", "Java", "SQL", "AWS"], exp=["5 years of experience in cloud engineering"],
         projects=["Cloud Cost Analyzer - Major Project: AWS cost dashboard"],
         certs=["AWS Certified Solutions Architect"]),
]


def build_lines(c):
    lines = [("title", c["name"]), ("text", f"{c['email']} | {c['phone']}"), ("text", f"Location: {c['location']}"),
             ("h", "Education"), ("text", c["edu"]), ("text", c["cgpa"]),
             ("h", "Technical Skills"), ("text", ", ".join(c["skills"])), ("h", "Experience")]
    lines += [("text", e) for e in c["exp"]]
    lines += [("h", "Projects")] + [("text", "• " + p) for p in c["projects"]]
    if c["certs"]:
        lines += [("h", "Certifications")] + [("text", "• " + x) for x in c["certs"]]
    return lines


def write_pdf(path, lines):
    ss = getSampleStyleSheet()
    story = []
    for kind, txt in lines:
        style = {"title": ss["Title"], "h": ss["Heading2"], "text": ss["BodyText"]}[kind]
        story += [Paragraph(txt.replace("&", "&amp;"), style), Spacer(1, 4)]
    SimpleDocTemplate(str(path), pagesize=A4).build(story)


def write_docx(path, lines):
    doc = Document()
    for kind, txt in lines:
        if kind == "title":
            doc.add_heading(txt, 0)
        elif kind == "h":
            doc.add_heading(txt, 1)
        else:
            doc.add_paragraph(txt)
    doc.save(str(path))


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    for c in CANDIDATES:
        lines = build_lines(c)
        path = OUT / c["file"]
        (write_pdf if c["file"].endswith(".pdf") else write_docx)(path, lines)
        print("created", path.name)


if __name__ == "__main__":
    main()
